"""
Export the live database back to a familiar Excel workbook.

Produces a file that mirrors Balaji's original sheets (Subscription Data,
Meals Tracker, Trial Orders, Customer Database, Monthly Dashboard, Finance
Sheet) so he always has a readable backup he can open in Excel — and a bridge
to feed the Finance / Money-OS later.
"""

import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

import db
import core

APP_DIR = os.path.dirname(os.path.abspath(__file__))
EXPORT_DIR = os.path.join(APP_DIR, "exports")

FOREST = "FF004030"
LIME = "FF1FA84A"
CREAM = "FFF8F5EC"


def _header(ws, headers, row=1):
    fill = PatternFill("solid", fgColor=FOREST)
    for ci, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=ci, value=h)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = f"A{row + 1}"   # string ref — avoids creating a blank data row


def export_workbook(path=None):
    os.makedirs(EXPORT_DIR, exist_ok=True)
    if not path:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        path = os.path.join(EXPORT_DIR, f"OG_Gainz_export_{stamp}.xlsx")

    c = db.get_db()
    wb = openpyxl.Workbook()

    # --- Subscription Data ---
    ws = wb.active
    ws.title = "SUBSCRIPTION DATA"
    _header(ws, ["Sub ID", "Cust ID", "Customer", "Plan", "Duration", "Notes", "Slot",
                 "Timing", "Booking", "Start", "Month", "Price", "Delivery Fee", "Total",
                 "Received", "Payment", "Scheduled", "Delivered", "Skipped", "Remaining",
                 "Status", "Last Delivery"])
    for r in c.execute("SELECT * FROM subscriptions ORDER BY id"):
        ws.append([r["id"], r["customer_id"], r["customer_name"], r["plan_name"], r["duration"],
                   r["notes"], r["slot"], r["preferred_timing"], r["booking_date"], r["start_date"],
                   r["month"], r["subscription_price"], r["delivery_fee"], r["total_amount"],
                   r["amount_received"], r["payment_status"], r["total_meals_scheduled"],
                   r["delivered_count"], r["skipped_count"], r["remaining_meals"], r["status"],
                   r["last_delivery_date"]])

    # --- Meals Tracker ---
    ws = wb.create_sheet("MEALS TRACKER")
    _header(ws, ["Date", "Sub ID", "Cust ID", "Name", "Meal Type", "Status", "Month", "Item/Notes"])
    for r in c.execute("SELECT * FROM deliveries ORDER BY date, name"):
        ws.append([r["date"], r["subscription_id"], r["customer_id"], r["name"],
                   r["meal_type"], r["status"], r["month"], r["notes"]])

    # --- Trial Orders ---
    ws = wb.create_sheet("TRIAL ORDERS")
    _header(ws, ["Order ID", "Date", "Name", "Phone", "Order", "Amount", "Payment",
                 "Month", "Conversion", "Address"])
    for r in c.execute("SELECT * FROM trial_orders ORDER BY id"):
        ws.append([r["id"], r["date"], r["name"], r["phone"], r["order_text"], r["amount"],
                   r["payment_status"], r["month"], r["conversion"], r["address"]])

    # --- Customer Database ---
    ws = wb.create_sheet("CUSTOMER DATABASE")
    _header(ws, ["Customer ID", "Name", "Phone", "Locality", "Address", "Notes"])
    for r in c.execute("SELECT * FROM customers ORDER BY id"):
        ws.append([r["id"], r["name"], r["phone"], r["locality"], r["address"], r["notes"]])

    # --- Monthly Dashboard ---
    ws = wb.create_sheet("MONTHLY DASHBOARD")
    _header(ws, ["Month", "Subscription Sales", "Trial Sales", "Overall Revenue",
                 "Subs Sold", "Paid", "Unpaid", "Orders Delivered", "Growth %"])
    for m in core.monthly_breakdown(c):
        ws.append([m["month"], m["sub_sales"], m["trial_sales"], m["overall"], m["subs_sold"],
                   m["paid"], m["unpaid"], m["delivered"],
                   "" if m["growth"] is None else round(m["growth"] * 100, 1)])

    # --- Finance Sheet ---
    ws = wb.create_sheet("FINANCE SHEET")
    s = core.headline_stats(c)
    ws["A1"] = "OVERALL FINANCIAL DASHBOARD"
    ws["A1"].font = Font(bold=True, size=14, color=FOREST)
    rows = [
        ("Total subscription sales", s["sales"]),
        ("Total revenue received", s["received"]),
        ("Pending amount", s["pending_amt"]),
        ("Total trial sales", s["trial_sales"]),
        ("OVERALL REVENUE", s["overall"]),
        ("Active subscriptions", s["active"]),
        ("Unpaid subscriptions", s["unpaid_n"]),
        ("Customers", s["customers"]),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)

    # widen columns a little
    for sheet in wb.worksheets:
        for col in sheet.columns:
            width = max((len(str(cell.value)) for cell in col if cell.value is not None), default=8)
            sheet.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 34)

    c.close()
    wb.save(path)
    return path


if __name__ == "__main__":
    print("Exported:", export_workbook())
