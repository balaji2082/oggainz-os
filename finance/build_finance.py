"""
OG Gainz — Finance Dashboard Builder
====================================

Pulls expenses from Notion + revenue from oggainz.db, then writes:
  * dashboard.html   — self-contained visual dashboard (opens in any browser, works offline)
  * OG-Gainz-Finance.xlsx — Excel workbook with monthly sheets

Run it with:  Update-Dashboard.bat   (or:  python finance/build_finance.py)

To EDIT entries from the dashboard instead of just reading them, use Edit-Dashboard.bat,
which runs finance/serve_dashboard.py — that's the only piece that writes to Notion.

SETUP — once only
-----------------
1. Go to  https://www.notion.so/my-integrations
2. "New integration"  -> name it "OG Gainz Finance" -> Submit
3. Copy the "Internal Integration Secret" (starts with  ntn_  or  secret_ )
4. Paste it into a file called  .notion_token  in the OG-Gainz-OS folder
5. IMPORTANT — open each database in Notion, click the ••• menu (top right),
   choose "Connections" -> "OG Gainz Finance". Without this the integration
   cannot see the database and you'll get a 404.

This file never writes to Notion or to the live app — it only reads. Writes happen in
serve_dashboard.py, and only when you have Edit-Dashboard.bat running.
"""

import json
import os
import sqlite3
import sys
import urllib.error
import urllib.request
from collections import defaultdict
from datetime import datetime

# Works whether this file is run as a script or imported from the app root.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dashboard_render import render_dashboard  # noqa: E402

# ---------------------------------------------------------------- paths / config

FIN_DIR = os.path.dirname(os.path.abspath(__file__))
APP_DIR = os.path.dirname(FIN_DIR)
DB_PATH = os.path.join(APP_DIR, "oggainz.db")
TOKEN_FILE = os.path.join(APP_DIR, ".notion_token")
DATA_DIR = os.path.join(FIN_DIR, "data")
OUT_HTML = os.path.join(FIN_DIR, "dashboard.html")
OUT_XLSX = os.path.join(FIN_DIR, "OG-Gainz-Finance.xlsx")

BUSINESS_DB = "9732da24eed54a20a97394e0a1d7ca77"
PERSONAL_DB = "b7de3b101c6e4a129a9a734c688015f8"
NOTION_VERSION = "2022-06-28"

# Brand
FOREST, LIME, ORANGE, CREAM, INK = "#004030", "#1FA84A", "#E9762B", "#F8F5EC", "#0A0A0A"
SAGE, PINE, LIMELIGHT = "#D6E5DB", "#0A6B3F", "#52C975"

# Zero-based costing model (One-Meal-Plan/02_Zero_Based_Costing.md)
MODEL_KITCHEN_GATE = {"Signature": 92.5, "Elite": 116.5, "Athlete": 146.0}
MODEL_PACKAGING_PER_MEAL = 19.0
MODEL_BLENDED_GATE = 112.0

MONTH_ORDER = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
               "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def month_key(m):
    """'Jul-2026' -> (2026, 7) for sorting."""
    if not m or "-" not in m:
        return (0, 0)
    mon, yr = m.split("-", 1)
    try:
        return (int(yr), MONTH_ORDER.index(mon[:3]) + 1)
    except (ValueError, IndexError):
        return (0, 0)


def iso_to_month(iso):
    """'2026-08-04' -> 'Aug-2026'."""
    if not iso:
        return None
    try:
        d = datetime.strptime(iso[:10], "%Y-%m-%d")
        return f"{MONTH_ORDER[d.month - 1]}-{d.year}"
    except ValueError:
        return None


def add_months(m, n):
    """'Aug-2026' + 3 -> 'Nov-2026'."""
    y, mo = month_key(m)
    if not y:
        return m
    total = (y * 12 + mo - 1) + n
    return f"{MONTH_ORDER[total % 12]}-{total // 12}"


# ---------------------------------------------------------------- notion

def get_token():
    tok = os.environ.get("OG_NOTION_TOKEN")
    if tok:
        return tok.strip()
    if os.path.exists(TOKEN_FILE):
        with open(TOKEN_FILE, encoding="utf-8") as f:
            t = f.read().strip()
            if t:
                return t
    return None


def notion_query(db_id, token):
    """Page through a Notion database. Returns list of raw page dicts."""
    url = f"https://api.notion.com/v1/databases/{db_id}/query"
    headers = {
        "Authorization": f"Bearer {token}",
        "Notion-Version": NOTION_VERSION,
        "Content-Type": "application/json",
    }
    rows, cursor = [], None
    while True:
        body = {"page_size": 100}
        if cursor:
            body["start_cursor"] = cursor
        req = urllib.request.Request(
            url, data=json.dumps(body).encode(), headers=headers, method="POST")
        try:
            with urllib.request.urlopen(req, timeout=60) as resp:
                payload = json.loads(resp.read())
        except urllib.error.HTTPError as e:
            detail = e.read().decode(errors="replace")[:400]
            if e.code == 404:
                raise SystemExit(
                    "\nNotion returned 404 for a database.\n"
                    "This almost always means the integration hasn't been given access.\n"
                    "Fix: open the database in Notion -> ••• menu (top right) ->\n"
                    "     Connections -> add your integration. Then run this again.\n")
            if e.code == 401:
                raise SystemExit(
                    f"\nNotion rejected the token (401).\nCheck {TOKEN_FILE} is correct.\n")
            raise SystemExit(f"\nNotion returned HTTP {e.code}:\n{detail}\n")
        except urllib.error.URLError as e:
            raise SystemExit(f"\nCould not reach Notion ({e.reason}). Check your internet.\n")
        rows.extend(payload.get("results", []))
        if not payload.get("has_more"):
            break
        cursor = payload.get("next_cursor")
    return rows


def prop_value(p):
    """Flatten one Notion property into a plain Python value."""
    if not isinstance(p, dict):
        return None
    t = p.get("type")
    if t == "title" or t == "rich_text":
        return "".join(x.get("plain_text", "") for x in p.get(t) or []) or None
    if t == "number":
        return p.get("number")
    if t == "select":
        return (p.get("select") or {}).get("name")
    if t == "date":
        return (p.get("date") or {}).get("start")
    if t == "files":
        return len(p.get("files") or []) or None
    if t == "formula":
        f = p.get("formula") or {}
        return f.get(f.get("type"))
    return None


def flatten(pages):
    out = []
    for pg in pages:
        row = {k: prop_value(v) for k, v in (pg.get("properties") or {}).items()}
        row["_id"] = pg.get("id")
        out.append(row)
    return out


def find_dad_db(token):
    """Locate appa's expense database, if it has been shared with the integration.

    Discovered by title rather than a hard-coded ID, so sharing it in Notion is the
    only step — nothing here needs editing. Returns (db_id, title) or (None, reason).
    """
    url = f"{'https://api.notion.com/v1'}/search"
    headers = {
        "Authorization": f"Bearer {token}",
        "Notion-Version": NOTION_VERSION,
        "Content-Type": "application/json",
    }
    body = {"filter": {"value": "database", "property": "object"}, "page_size": 100}
    req = urllib.request.Request(
        url, data=json.dumps(body).encode(), headers=headers, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            payload = json.loads(resp.read())
    except (urllib.error.HTTPError, urllib.error.URLError):
        return None, "could not search Notion"

    known = {BUSINESS_DB, PERSONAL_DB}
    hints = ("appa", "dad", "father", "papa")
    fallback = None
    for db in payload.get("results", []):
        db_id = (db.get("id") or "").replace("-", "")
        if db_id in known:
            continue
        title = "".join(x.get("plain_text", "") for x in db.get("title", []))
        low = title.lower()
        if any(h in low for h in hints):
            return db_id, title
        if "expense" in low or "spend" in low:
            fallback = (db_id, title)
    if fallback:
        return fallback
    return None, "not shared with the integration"


def shape_shared(rows):
    """Tolerant shaper for appa's database.

    Its schema isn't ours to control, so match property names case-insensitively and
    accept the common variants rather than assuming. Anything missing is left blank —
    the dashboard renders fine without it.
    """
    def pick(row, *names):
        low = {k.lower(): v for k, v in row.items()}
        for n in names:
            v = low.get(n.lower())
            if v not in (None, ""):
                return v
        return None

    def yes_no(v):
        """Normalise a Yes/No-ish flag to True / False / None."""
        if v is None:
            return None
        if isinstance(v, bool):
            return v
        s = str(v).strip().lower()
        if s in ("yes", "y", "true", "business", "1"):
            return True
        if s in ("no", "n", "false", "personal", "home", "0"):
            return False
        return None

    clean = []
    for r in rows:
        amt = pick(r, "Amount", "Cost", "Price", "Value", "Spent")
        try:
            amt = float(amt)
        except (TypeError, ValueError):
            continue
        if not amt:
            continue
        date = pick(r, "Date", "Day", "When", "Spent On")
        if not date:
            continue
        cat = pick(r, "Category", "Type", "Head") or "Uncategorised"
        # The "did this feed the kitchen?" flag, however he named it. Appa buys for
        # both the business and the house out of one pocket, so this is the field
        # that decides whether a row is really an OG Gainz cost.
        to_biz = yes_no(pick(r, "Goes To Business", "For Business", "Business?",
                             "Is Business", "For", "Side", "Business", "Purpose"))
        clean.append({
            "id": r.get("_id"), "db": "dad",
            "date": str(date)[:10],
            "month": iso_to_month(str(date)),
            "what": pick(r, "What", "Name", "Item", "Title", "Description"),
            "amount": amt,
            "category": cat,
            "cost_type": pick(r, "Cost Type") or None,
            "subcategory": pick(r, "Subcategory", "Sub Category", "Sub"),
            "vendor": pick(r, "Vendor", "Shop", "Store", "Merchant", "Paid To"),
            "payment": pick(r, "Payment Method", "Paid By", "Mode", "Payment"),
            "to_business": to_biz,
            "notes": pick(r, "Notes", "Note", "Remarks", "Comment"),
        })
    clean.sort(key=lambda x: x["date"] or "", reverse=True)
    return clean


# ---------------------------------------------------------------- revenue (sqlite)

def load_revenue():
    if not os.path.exists(DB_PATH):
        raise SystemExit(f"\nCouldn't find {DB_PATH}.\nRun Get-Live-Data.bat first.\n")
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    rev = defaultdict(lambda: {
        "food": 0.0, "delivery": 0.0, "trial": 0.0, "billed": 0.0, "received": 0.0,
        "subs": 0, "delivered": 0, "skipped": 0, "pending": 0})

    for r in c.execute("""SELECT month, COUNT(*) n, SUM(subscription_price) sp,
                          SUM(delivery_fee) df, SUM(total_amount) ta,
                          SUM(amount_received) ar FROM subscriptions GROUP BY month"""):
        if not r["month"]:
            continue
        m = rev[r["month"]]
        m["food"] = r["sp"] or 0
        m["delivery"] = r["df"] or 0
        m["billed"] = r["ta"] or 0
        m["received"] = r["ar"] or 0
        m["subs"] = r["n"]

    for r in c.execute("SELECT month, SUM(amount) a FROM trial_orders GROUP BY month"):
        if r["month"]:
            rev[r["month"]]["trial"] = r["a"] or 0

    for r in c.execute("""SELECT month, status, COUNT(*) n FROM deliveries
                          GROUP BY month, status"""):
        if not r["month"]:
            continue
        s = (r["status"] or "").lower()
        if s == "delivered":
            rev[r["month"]]["delivered"] = r["n"]
        elif s == "skipped":
            rev[r["month"]]["skipped"] = r["n"]
        elif s == "pending":
            rev[r["month"]]["pending"] = r["n"]

    # --- Recognised revenue: value the meals actually DELIVERED in each month.
    # Subscriptions are billed upfront but delivered over weeks, so raw monthly
    # billing divided by that month's meals is meaningless. Each delivered meal is
    # worth (its subscription's food price / that subscription's total meals).
    for r in c.execute("""
        SELECT d.month AS m,
               SUM(CASE WHEN s.total_meals_scheduled > 0
                        THEN s.subscription_price * 1.0 / s.total_meals_scheduled
                        ELSE 0 END) AS val
        FROM deliveries d
        JOIN subscriptions s ON s.id = d.subscription_id
        WHERE d.status = 'Delivered'
        GROUP BY d.month"""):
        if r["m"]:
            rev[r["m"]]["recognised"] = r["val"] or 0.0

    plans = {r["name"]: dict(r) for r in c.execute("SELECT * FROM plans")}
    c.close()

    for m in rev.values():
        m["total"] = m["billed"] + m["trial"]
        m["outstanding"] = m["billed"] - m["received"]
        m.setdefault("recognised", 0.0)
    return dict(rev), plans


# ---------------------------------------------------------------- expense shaping

def shape_expenses(rows):
    """Split each expense into cash-in-month and amortised monthly slices."""
    cash = defaultdict(lambda: defaultdict(float))      # month -> cost_type -> amount
    cash_cat = defaultdict(lambda: defaultdict(float))  # month -> category -> amount
    true_cost = defaultdict(lambda: defaultdict(float))  # month -> cost_type -> amortised
    delivery = defaultdict(lambda: {"spend": 0.0, "orders": 0, "runs": 0})
    vendors = defaultdict(float)
    clean = []

    for r in rows:
        amt = r.get("Amount") or 0
        if not amt:
            continue
        m = iso_to_month(r.get("Date"))
        if not m:
            continue
        cat = r.get("Category") or "Uncategorised"
        ctype = r.get("Cost Type") or (cat.split(" · ")[0] if " · " in cat else "OPEX")
        spread = r.get("Spread Over Months") or 1
        try:
            spread = max(1.0, float(spread))
        except (TypeError, ValueError):
            spread = 1.0

        cash[m][ctype] += amt
        cash_cat[m][cat] += amt
        if r.get("Vendor"):
            vendors[r["Vendor"]] += amt

        # amortise across months
        slice_amt = amt / spread
        whole = int(spread)
        for i in range(whole):
            true_cost[add_months(m, i)][ctype] += slice_amt
        rem = spread - whole
        if rem > 0.01:
            true_cost[add_months(m, whole)][ctype] += slice_amt * rem

        if cat == "COGS · Delivery":
            d = delivery[m]
            d["spend"] += amt
            d["runs"] += 1
            d["orders"] += int(r.get("Orders Covered") or 0)

        clean.append({
            # id + db let the dashboard write a correction straight back to Notion
            "id": r.get("_id"), "db": "business",
            "month": m, "date": r.get("Date"), "what": r.get("What"),
            "amount": amt, "category": cat, "cost_type": ctype,
            "subcategory": r.get("Subcategory"), "vendor": r.get("Vendor"),
            "spread": spread, "true_monthly": round(slice_amt),
            "payment": r.get("Payment Method"), "paid_from": r.get("Paid From"),
            "orders_covered": r.get("Orders Covered"), "notes": r.get("Notes"),
        })

    clean.sort(key=lambda x: x["date"] or "", reverse=True)
    return clean, dict(cash), dict(cash_cat), dict(true_cost), dict(delivery), dict(vendors)


def shape_personal(rows):
    months = defaultdict(lambda: {"out": 0.0, "in": 0.0})
    cats = defaultdict(lambda: defaultdict(float))
    clean = []
    for r in rows:
        amt = r.get("Amount") or 0
        m = iso_to_month(r.get("Date"))
        if not amt or not m:
            continue
        direction = r.get("Direction") or "Money Out"
        cat = r.get("Category") or "Miscellaneous"
        if direction == "Money In":
            months[m]["in"] += amt
        else:
            months[m]["out"] += amt
            cats[m][cat] += amt
        clean.append({"id": r.get("_id"), "db": "personal",
                      "month": m, "date": r.get("Date"), "what": r.get("What"),
                      "amount": amt, "category": cat, "direction": direction,
                      "subcategory": r.get("Subcategory"),
                      "payment": r.get("Payment Method"), "notes": r.get("Notes")})
    clean.sort(key=lambda x: x["date"] or "", reverse=True)
    return clean, dict(months), dict(cats)


def build_monthly(rev, cash, true_cost, delivery, personal_m):
    """One row per month with everything joined and derived."""
    months = sorted(set(rev) | set(cash) | set(true_cost) | set(personal_m), key=month_key)
    out = []
    for m in months:
        r = rev.get(m, {})
        c = cash.get(m, {})
        t = true_cost.get(m, {})
        d = delivery.get(m, {})
        p = personal_m.get(m, {})

        cogs_true = t.get("COGS", 0.0)
        opex_true = t.get("OPEX", 0.0)
        capex_cash = c.get("CAPEX", 0.0)
        draw = c.get("DRAW", 0.0)
        delivered = r.get("delivered", 0)
        food = r.get("food", 0.0)
        recognised = r.get("recognised", 0.0)
        deliv_rev = r.get("delivery", 0.0)

        # Margin is measured against RECOGNISED revenue (value of meals actually
        # delivered this month), never against billing — see DATA-AUDIT.md §7.
        # Margin is meaningless until costs are actually logged — showing 100%
        # because nothing was entered would be worse than showing nothing.
        has_costs = cogs_true > 0
        gross = (recognised - cogs_true) if (recognised and has_costs) else 0.0
        row = {
            "month": m,
            "food_revenue": food,
            "recognised_revenue": recognised,
            "revenue_per_meal": (recognised / delivered) if delivered else None,
            "delivery_revenue": deliv_rev,
            "trial_revenue": r.get("trial", 0.0),
            "total_revenue": r.get("total", 0.0),
            "received": r.get("received", 0.0),
            "outstanding": r.get("outstanding", 0.0),
            "subs": r.get("subs", 0),
            "delivered": delivered,
            "skipped": r.get("skipped", 0),
            "pending": r.get("pending", 0),
            "cogs_cash": c.get("COGS", 0.0),
            "cogs_true": cogs_true,
            "opex_true": opex_true,
            "capex": capex_cash,
            "draw": draw,
            "cash_out": sum(c.values()),
            "gross_profit": gross if has_costs else None,
            "gross_margin": (gross / recognised * 100) if (recognised and has_costs) else None,
            "net_profit": (recognised + deliv_rev + r.get("trial", 0.0)
                           - cogs_true - opex_true) if (recognised and has_costs) else None,
            "cost_per_meal": (cogs_true / delivered) if delivered and cogs_true else None,
            "delivery_spend": d.get("spend", 0.0),
            "delivery_orders": d.get("orders", 0),
            "delivery_runs": d.get("runs", 0),
            "delivery_net": deliv_rev - d.get("spend", 0.0),
            "cost_per_delivery": (d["spend"] / d["orders"]) if d.get("orders") else None,
            "skip_rate": (r.get("skipped", 0) / (delivered + r.get("skipped", 0)) * 100)
                         if (delivered + r.get("skipped", 0)) else None,
            "personal_out": p.get("out", 0.0),
            "personal_in": p.get("in", 0.0),
        }
        out.append(row)
    return out


# ---------------------------------------------------------------- svg charts

def rupee(v, dp=0):
    if v is None:
        return "—"
    return f"₹{v:,.{dp}f}"


def esc(s):
    return (str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def grouped_bars(months, series, width=760, height=260):
    """series = [(label, colour, [values...]), ...]"""
    if not months:
        return "<p class='muted'>No data yet.</p>"
    pad_l, pad_b, pad_t = 62, 30, 14
    plot_w, plot_h = width - pad_l - 12, height - pad_b - pad_t
    peak = max([v for _, _, vals in series for v in vals if v] or [1])
    peak = peak * 1.15 or 1
    gw = plot_w / len(months)
    bw = min(26, gw / (len(series) + 1))

    p = [f'<svg viewBox="0 0 {width} {height}" class="chart">']
    for i in range(5):  # gridlines
        y = pad_t + plot_h * i / 4
        val = peak * (1 - i / 4)
        p.append(f'<line x1="{pad_l}" y1="{y:.0f}" x2="{width-12}" y2="{y:.0f}" '
                 f'stroke="#e6ece8" stroke-width="1"/>')
        p.append(f'<text x="{pad_l-8}" y="{y+4:.0f}" class="ax" text-anchor="end">'
                 f'{val/1000:.0f}k</text>')
    for gi, m in enumerate(months):
        gx = pad_l + gi * gw
        for si, (label, colour, vals) in enumerate(series):
            v = vals[gi] or 0
            h = (v / peak) * plot_h if peak else 0
            x = gx + (gw - bw * len(series)) / 2 + si * bw
            y = pad_t + plot_h - h
            p.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{bw-3:.1f}" height="{max(h,0):.1f}" '
                     f'fill="{colour}" rx="2"><title>{esc(label)} {m}: {rupee(v)}</title></rect>')
        p.append(f'<text x="{gx+gw/2:.0f}" y="{height-10}" class="ax" '
                 f'text-anchor="middle">{esc(m[:3])}</text>')
    p.append("</svg>")
    legend = " ".join(
        f'<span class="key"><i style="background:{c}"></i>{esc(l)}</span>'
        for l, c, _ in series)
    return "".join(p) + f'<div class="legend">{legend}</div>'


def donut(pairs, size=210):
    """pairs = [(label, value, colour)]"""
    total = sum(v for _, v, _ in pairs)
    if not total:
        return "<p class='muted'>No spending logged yet.</p>"
    cx = cy = size / 2
    r, thick = size / 2 - 10, 30
    import math
    angle = -math.pi / 2
    p = [f'<svg viewBox="0 0 {size} {size}" class="donut">']
    for label, val, colour in pairs:
        if val <= 0:
            continue
        sweep = 2 * math.pi * (val / total)
        x1, y1 = cx + r * math.cos(angle), cy + r * math.sin(angle)
        angle += sweep
        x2, y2 = cx + r * math.cos(angle), cy + r * math.sin(angle)
        large = 1 if sweep > math.pi else 0
        p.append(f'<path d="M {x1:.2f} {y1:.2f} A {r} {r} 0 {large} 1 {x2:.2f} {y2:.2f}" '
                 f'fill="none" stroke="{colour}" stroke-width="{thick}">'
                 f'<title>{esc(label)}: {rupee(val)} ({val/total*100:.0f}%)</title></path>')
    p.append(f'<text x="{cx}" y="{cy-2}" text-anchor="middle" class="dnum">{rupee(total)}</text>')
    p.append(f'<text x="{cx}" y="{cy+16}" text-anchor="middle" class="dlab">total</text>')
    p.append("</svg>")
    legend = "".join(
        f'<div class="lrow"><i style="background:{c}"></i>'
        f'<span>{esc(l)}</span><b>{rupee(v)}</b></div>'
        for l, v, c in sorted(pairs, key=lambda x: -x[1]) if v > 0)
    return f'<div class="donutwrap">{"".join(p)}<div class="llist">{legend}</div></div>'


CSS = """
*{box-sizing:border-box;margin:0;padding:0}
body{font:15px/1.5 'Segoe UI',Roboto,'Trebuchet MS',sans-serif;background:#F8F5EC;color:#0A0A0A;padding:26px}
.wrap{max-width:1180px;margin:0 auto}
header{background:linear-gradient(135deg,#52C975,#0A6B3F);color:#fff;padding:24px 28px;border-radius:14px;margin-bottom:22px}
header h1{font-size:26px;letter-spacing:-.5px}
header p{opacity:.9;font-size:13px;margin-top:4px}
h2{font-size:15px;text-transform:uppercase;letter-spacing:1.2px;color:#004030;margin:26px 0 12px;font-weight:700}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(168px,1fr));gap:12px}
.card{background:#fff;border-radius:11px;padding:15px 17px;border:1px solid #e2e9e4}
.card .lab{font-size:11px;text-transform:uppercase;letter-spacing:.8px;color:#6b7d72;font-weight:600}
.card .val{font-size:24px;font-weight:700;color:#004030;margin-top:5px;letter-spacing:-.5px}
.card .sub{font-size:11.5px;color:#7b8a80;margin-top:3px}
.card.warn{border-color:#E9762B;background:#fff8f3}.card.warn .val{color:#E9762B}
.card.good{border-color:#1FA84A}.card.good .val{color:#1FA84A}
.panel{background:#fff;border-radius:12px;padding:19px 21px;border:1px solid #e2e9e4;margin-bottom:14px}
.grid2{display:grid;grid-template-columns:1.35fr 1fr;gap:14px}
@media(max-width:900px){.grid2{grid-template-columns:1fr}}
table{width:100%;border-collapse:collapse;font-size:13.5px}
th{text-align:left;font-size:10.5px;text-transform:uppercase;letter-spacing:.7px;color:#6b7d72;padding:7px 9px;border-bottom:2px solid #e2e9e4}
td{padding:7px 9px;border-bottom:1px solid #f0f4f1}
td.n,th.n{text-align:right;font-variant-numeric:tabular-nums}
tr:hover td{background:#fafcfa}
.chart{width:100%;height:auto}
.ax{font-size:10px;fill:#8a9a90}
.legend{display:flex;gap:15px;flex-wrap:wrap;margin-top:9px;font-size:12px;color:#5d6d64}
.key i,.lrow i{display:inline-block;width:10px;height:10px;border-radius:2px;margin-right:6px;vertical-align:-1px}
.donutwrap{display:flex;gap:20px;align-items:center;flex-wrap:wrap}
.donut{width:200px;flex:none}
.dnum{font-size:19px;font-weight:700;fill:#004030}.dlab{font-size:10px;fill:#8a9a90}
.llist{flex:1;min-width:180px}
.lrow{display:flex;align-items:center;gap:5px;padding:3px 0;font-size:12.5px;border-bottom:1px solid #f4f7f5}
.lrow span{flex:1}.lrow b{font-variant-numeric:tabular-nums}
.muted{color:#8a9a90;font-size:13px;padding:14px 0}
.note{background:#F8F5EC;border-left:3px solid #E9762B;padding:11px 14px;border-radius:0 7px 7px 0;font-size:13px;margin-top:12px;color:#4a5a50}
.pill{display:inline-block;padding:2px 8px;border-radius:20px;font-size:10.5px;font-weight:600}
.p-cogs{background:#fde8e8;color:#a33}.p-opex{background:#e5eefb;color:#2a5599}
.p-capex{background:#f0e8fb;color:#6b3fa0}.p-draw{background:#f2ede6;color:#7a6a52}
.split{display:grid;grid-template-columns:1fr 1fr;gap:14px}
@media(max-width:700px){.split{grid-template-columns:1fr}}
.big{font-size:31px;font-weight:700;color:#004030;letter-spacing:-1px}
footer{text-align:center;color:#9aa8a0;font-size:11.5px;margin-top:30px;padding-top:16px;border-top:1px solid #e2e9e4}
"""


def render_html(monthly, expenses, personal, vendors, cur):
    """cur = the month to feature (latest with data)."""
    row = next((m for m in monthly if m["month"] == cur), {})
    prev_m = None
    for i, m in enumerate(monthly):
        if m["month"] == cur and i > 0:
            prev_m = monthly[i - 1]
    months = [m["month"] for m in monthly][-8:]
    sub = [m for m in monthly if m["month"] in months]

    def card(lab, val, s="", cls=""):
        return (f'<div class="card {cls}"><div class="lab">{esc(lab)}</div>'
                f'<div class="val">{val}</div>'
                f'<div class="sub">{esc(s)}</div></div>')

    # ---- headline
    gm = row.get("gross_margin")
    cpm = row.get("cost_per_meal")
    has_exp = row.get("cash_out", 0) > 0
    cards = [
        card("Revenue billed", rupee(row.get("total_revenue")),
             f'{row.get("subs",0)} subscriptions'),
        card("Money received", rupee(row.get("received")),
             f'{rupee(row.get("outstanding"))} still owed',
             "warn" if row.get("outstanding", 0) > 0 else "good"),
        card("Meals delivered", f'{row.get("delivered",0):,}',
             f'{row.get("skipped",0)} skipped'
             + (f' · {row["skip_rate"]:.0f}%' if row.get("skip_rate") else "")),
        card("Spent (cash out)", rupee(row.get("cash_out")) if has_exp else "—",
             "logged in Notion" if has_exp else "start 1 Aug"),
        card("Gross margin", f'{gm:.0f}%' if gm else "—",
             "model says 51–54%", "good" if gm and gm >= 45 else ""),
        card("Cost per meal", rupee(cpm) if cpm else "—",
             f'model ~{rupee(MODEL_BLENDED_GATE)}'),
        card("Earned this month", rupee(row.get("recognised_revenue")),
             f'{rupee(row.get("revenue_per_meal"))}/meal delivered'),
    ]

    # ---- model vs actual
    model_rows = ""
    if has_exp and row.get("delivered"):
        pack = next((e for e in expenses
                     if e["month"] == cur and e["category"] == "COGS · Packaging"), None)
        pack_total = sum(e["true_monthly"] for e in expenses
                         if e["month"] == cur and e["category"] == "COGS · Packaging")
        checks = [
            ("Kitchen-gate cost per meal", MODEL_BLENDED_GATE, cpm),
            ("Packaging per meal", MODEL_PACKAGING_PER_MEAL,
             pack_total / row["delivered"] if row["delivered"] else None),
            ("Gross margin %", 53.0, gm),
        ]
        for label, model, actual in checks:
            if actual is None:
                model_rows += (f'<tr><td>{esc(label)}</td><td class="n">{model:,.1f}</td>'
                               f'<td class="n">—</td><td class="n">—</td></tr>')
                continue
            diff = actual - model
            colour = LIME if abs(diff) / model < 0.1 else ORANGE
            model_rows += (
                f'<tr><td>{esc(label)}</td><td class="n">{model:,.1f}</td>'
                f'<td class="n"><b>{actual:,.1f}</b></td>'
                f'<td class="n" style="color:{colour}">{diff:+,.1f}</td></tr>')
        model_panel = (
            '<table><tr><th>Line</th><th class="n">Model says</th>'
            '<th class="n">Actual</th><th class="n">Gap</th></tr>'
            f'{model_rows}</table>'
            '<div class="note">Model comes from <b>One-Meal-Plan/02_Zero_Based_Costing.md</b>. '
            'A gap here means the estimate needs replacing with your real number — '
            'that is the whole point of tracking.</div>')
    else:
        model_panel = ('<p class="muted">Needs one month of logged expenses. '
                       'From 1 August this compares your real costs against the '
                       'zero-based costing model.</p>')

    # ---- delivery P&L
    dspend, dorders = row.get("delivery_spend", 0), row.get("delivery_orders", 0)
    if dspend:
        per = row.get("cost_per_delivery")
        dnet = row.get("delivery_net", 0)
        delivery_panel = f"""
        <div class="cards">
          {card("Fees collected", rupee(row.get("delivery_revenue")))}
          {card("Porter + fuel paid", rupee(dspend), f'{row.get("delivery_runs",0)} runs')}
          {card("Net position", rupee(dnet), "", "good" if dnet >= 0 else "warn")}
          {card("Cost per order", rupee(per) if per else "—",
                f'{dorders} orders across {row.get("delivery_runs",0)} runs')}
        </div>
        <div class="note">Batched self-delivery costs ~₹2–5/serving. A solo Porter run to Z3
        costs ₹120–130 against a ₹60–75 fee. <b>Cost per order</b> is the number that proves
        whether batching is working — log <i>Orders Covered</i> on every Porter run.</div>"""
    else:
        delivery_panel = ('<p class="muted">No delivery costs logged yet. Log Porter runs under '
                          '<b>COGS · Delivery</b> and fill <b>Orders Covered</b> — that is what '
                          'shows whether batching saves money.</p>')

    # ---- charts
    rev_chart = grouped_bars(months, [
        ("Revenue billed", LIME, [m["total_revenue"] for m in sub]),
        ("Received", PINE, [m["received"] for m in sub]),
        ("Spent", ORANGE, [m["cash_out"] for m in sub]),
    ])
    cost_chart = grouped_bars(months, [
        ("COGS", ORANGE, [m["cogs_true"] for m in sub]),
        ("Operating", FOREST, [m["opex_true"] for m in sub]),
        ("Capex", "#6b3fa0", [m["capex"] for m in sub]),
    ], height=210)

    cat_totals = defaultdict(float)
    for e in expenses:
        if e["month"] == cur:
            cat_totals[e["category"]] += e["amount"]
    palette = [ORANGE, LIME, FOREST, "#6b3fa0", LIMELIGHT, "#a33", "#2a5599", "#7a6a52",
               "#d4a017", SAGE]
    cat_donut = donut([(k.split(" · ")[-1], v, palette[i % len(palette)])
                       for i, (k, v) in enumerate(sorted(cat_totals.items(),
                                                         key=lambda x: -x[1]))])

    # ---- tables
    exp_rows = ""
    for e in expenses[:40]:
        cls = "p-" + (e["cost_type"] or "opex").lower()
        spread_note = (f' <span class="muted">÷{e["spread"]:g}</span>'
                       if e["spread"] and e["spread"] > 1 else "")
        exp_rows += (
            f'<tr><td>{esc(e["date"])}</td><td>{esc(e["what"] or e["subcategory"] or "")}</td>'
            f'<td><span class="pill {cls}">{esc(e["cost_type"])}</span> '
            f'{esc(e["category"].split(" · ")[-1])}</td>'
            f'<td class="n">{rupee(e["amount"])}{spread_note}</td>'
            f'<td class="n">{rupee(e["true_monthly"])}</td>'
            f'<td>{esc(e["paid_from"] or "")}</td></tr>')
    if not exp_rows:
        exp_rows = '<tr><td colspan="6" class="muted">Nothing logged yet — starts 1 August.</td></tr>'

    # Say it plainly when the featured month is lopsided, instead of quietly
    # showing ₹0 and leaving you to wonder whether it's broken.
    gap_note = ""
    if row.get("total_revenue") and not row.get("cash_out"):
        gap_note = (f'<div class="note"><b>{esc(cur)} has revenue but no logged expenses.</b> '
                    'Everything cost-related below reads ₹0 for that reason — nothing is '
                    'broken. Log an expense dated in this month and it will populate.</div>')
    elif row.get("cash_out") and not row.get("total_revenue"):
        gap_note = (f'<div class="note"><b>{esc(cur)} has expenses but no revenue yet.</b> '
                    'Margin and cost per meal stay blank until subscriptions and deliveries '
                    'land in this month. Expected if the month has only just started.</div>')

    today_m = f"{MONTH_ORDER[datetime.now().month - 1]}-{datetime.now().year}"
    mon_rows = ""
    future_seen = False
    for m in reversed(monthly):
        margin_txt = f'{m["gross_margin"]:.0f}%' if m["gross_margin"] else "—"
        cogs_txt = rupee(m["cogs_true"]) if m["cogs_true"] else "—"
        opex_txt = rupee(m["opex_true"]) if m["opex_true"] else "—"
        cpm_txt = rupee(m["cost_per_meal"], 1) if m["cost_per_meal"] else "—"
        is_future = month_key(m["month"]) > month_key(today_m)
        if is_future:
            future_seen = True
            cogs_txt += ' <span class="muted">committed</span>'
        mon_rows += (
            f'<tr{" style=opacity:.55" if is_future else ""}>'
            f'<td><b>{esc(m["month"])}</b></td>'
            f'<td class="n">{rupee(m["total_revenue"])}</td>'
            f'<td class="n">{rupee(m["received"])}</td>'
            f'<td class="n">{m["delivered"]:,}</td>'
            f'<td class="n">{cogs_txt}</td>'
            f'<td class="n">{opex_txt}</td>'
            f'<td class="n">{margin_txt}</td>'
            f'<td class="n">{cpm_txt}</td></tr>')

    # ---- personal
    prow = row
    pcats = defaultdict(float)
    for p in personal:
        if p["month"] == cur and p["direction"] != "Money In":
            pcats[p["category"]] += p["amount"]
    p_donut = donut([(k, v, palette[i % len(palette)])
                     for i, (k, v) in enumerate(sorted(pcats.items(), key=lambda x: -x[1]))])
    draw = prow.get("draw", 0) or prow.get("personal_in", 0)
    pout = prow.get("personal_out", 0)

    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>OG Gainz — Finance</title><style>{CSS}</style></head><body><div class="wrap">
<header>
  <h1>OG Gainz — Finance</h1>
  <p>Featuring <b>{esc(cur)}</b> · generated {datetime.now():%d %b %Y, %H:%M} ·
  revenue from the live app, expenses from Notion</p>
</header>

<h2>This month</h2>
<div class="cards">{''.join(cards)}</div>
{gap_note}
<div class="note"><b>Billed vs earned.</b> Subscriptions are paid upfront but delivered over
weeks, so a month's billing does not match the food it cooked. <b>Revenue billed</b> is what
you invoiced; <b>earned this month</b> values only the meals actually delivered. Margin and
cost per meal are measured against <i>earned</i> — comparing costs to billing would make
good months look bad and bad months look great.</div>

<div class="grid2" style="margin-top:14px">
  <div class="panel"><h2 style="margin-top:0">Revenue vs spending</h2>{rev_chart}</div>
  <div class="panel"><h2 style="margin-top:0">Where the money went</h2>{cat_donut}</div>
</div>

<div class="panel"><h2 style="margin-top:0">Model vs actual</h2>{model_panel}</div>

<div class="panel"><h2 style="margin-top:0">Delivery P&amp;L</h2>{delivery_panel}</div>

<div class="panel"><h2 style="margin-top:0">Cost structure over time</h2>{cost_chart}</div>

<h2>Month by month</h2>
<div class="panel" style="overflow-x:auto">
<table><tr><th>Month</th><th class="n">Billed</th><th class="n">Received</th>
<th class="n">Meals</th><th class="n">COGS</th><th class="n">Operating</th>
<th class="n">Margin</th><th class="n">Cost/meal</th></tr>{mon_rows}</table>
{'<div class="note"><b>Faded rows are future months.</b> They already carry cost because a bulk buy was spread forward — those trays are paid for and August will eat them. Revenue is zero simply because it has not happened yet.</div>' if future_seen else ''}
</div>

<h2>Recent expenses</h2>
<div class="panel" style="overflow-x:auto">
<table><tr><th>Date</th><th>What</th><th>Category</th><th class="n">Paid</th>
<th class="n">True monthly</th><th>From</th></tr>{exp_rows}</table>
</div>

<h2>Personal — kept separate</h2>
<div class="split">
  <div class="panel">
    <div class="lab" style="font-size:11px;color:#6b7d72;text-transform:uppercase;letter-spacing:.8px">My personal spend</div>
    <div class="big">{rupee(pout)}</div>
    <div class="sub" style="font-size:12px;color:#7b8a80">
      Owner draw taken: {rupee(draw)}</div>
    {'<div class="note">You spent more personally than you drew from the business this month.</div>' if pout > draw > 0 else ''}
  </div>
  <div class="panel">
    <div class="lab" style="font-size:11px;color:#6b7d72;text-transform:uppercase;letter-spacing:.8px">OG Gainz spend</div>
    <div class="big">{rupee(row.get('cash_out'))}</div>
    <div class="sub" style="font-size:12px;color:#7b8a80">Business money — deliberately not added to your personal total</div>
  </div>
</div>
<div class="panel" style="margin-top:14px"><h2 style="margin-top:0">Personal categories</h2>{p_donut}</div>

<footer>Generated by <b>build_finance.py</b> · run <b>Update-Dashboard.bat</b> to refresh<br>
Run <b>Get-Live-Data.bat</b> first if revenue looks out of date</footer>
</div></body></html>"""


# ---------------------------------------------------------------- excel

def write_excel(monthly, expenses, personal, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("  ! openpyxl not installed — skipping Excel. Run: pip install openpyxl")
        return False

    hdr_fill = PatternFill("solid", fgColor="004030")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    money = '₹#,##0'

    wb = Workbook()

    def sheet(title, headers, rows, widths, money_cols=()):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(1, c)
            cell.fill, cell.font = hdr_fill, hdr_font
            cell.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        for col in money_cols:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, col).number_format = money
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)

    sheet("Monthly Summary",
          ["Month", "Revenue billed", "Earned (delivered)", "Received", "Outstanding",
           "Subs", "Meals delivered", "Skipped", "Revenue per meal", "COGS (true)",
           "Operating", "Capex", "Owner draw", "Cash out", "Gross profit",
           "Gross margin %", "Cost per meal", "Delivery net", "Personal spend"],
          [[m["month"], m["total_revenue"], m["recognised_revenue"], m["received"],
            m["outstanding"], m["subs"], m["delivered"], m["skipped"],
            round(m["revenue_per_meal"], 1) if m["revenue_per_meal"] else None,
            m["cogs_true"], m["opex_true"], m["capex"], m["draw"], m["cash_out"],
            m["gross_profit"],
            round(m["gross_margin"], 1) if m["gross_margin"] else None,
            round(m["cost_per_meal"], 1) if m["cost_per_meal"] else None,
            m["delivery_net"], m["personal_out"]] for m in monthly],
          [11, 15, 17, 13, 13, 7, 15, 9, 15, 13, 12, 10, 12, 12, 13, 14, 13, 13, 14],
          money_cols=(2, 3, 4, 5, 9, 10, 11, 12, 13, 14, 15, 17, 18, 19))

    sheet("Business Expenses",
          ["Date", "Month", "What", "Cost Type", "Category", "Subcategory", "Amount",
           "Spread", "True monthly", "Vendor", "Payment", "Paid from", "Orders covered",
           "Notes"],
          [[e["date"], e["month"], e["what"], e["cost_type"], e["category"],
            e["subcategory"], e["amount"], e["spread"], e["true_monthly"], e["vendor"],
            e["payment"], e["paid_from"], e["orders_covered"], e["notes"]]
           for e in expenses],
          [11, 10, 26, 10, 24, 17, 11, 8, 13, 18, 12, 14, 14, 30],
          money_cols=(7, 9))

    sheet("Personal Expenses",
          ["Date", "Month", "What", "Direction", "Category", "Subcategory", "Amount",
           "Payment", "Notes"],
          [[p["date"], p["month"], p["what"], p["direction"], p["category"],
            p["subcategory"], p["amount"], p["payment"], p["notes"]] for p in personal],
          [11, 10, 26, 11, 22, 18, 11, 12, 30],
          money_cols=(7,))

    sheet("Revenue Detail",
          ["Month", "Food revenue", "Delivery fees", "Trials", "Total billed", "Received",
           "Outstanding", "Subscriptions", "Delivered", "Skipped", "Pending", "Skip rate %"],
          [[m["month"], m["food_revenue"], m["delivery_revenue"], m["trial_revenue"],
            m["total_revenue"], m["received"], m["outstanding"], m["subs"], m["delivered"],
            m["skipped"], m["pending"],
            round(m["skip_rate"], 1) if m["skip_rate"] else None] for m in monthly],
          [11, 14, 14, 11, 13, 13, 13, 14, 11, 10, 10, 11],
          money_cols=(2, 3, 4, 5, 6, 7))

    wb.save(path)
    return True


# ---------------------------------------------------------------- main

def main():
    print("\n  OG GAINZ — FINANCE DASHBOARD\n  " + "-" * 44)

    print("  Reading revenue from oggainz.db ...")
    rev, plans = load_revenue()
    print(f"    {len(rev)} months of revenue found")

    token = get_token()
    biz_rows, per_rows, dad_rows = [], [], []
    dad_status = None
    if token:
        print("  Reading expenses from Notion ...")
        biz_rows = flatten(notion_query(BUSINESS_DB, token))
        per_rows = flatten(notion_query(PERSONAL_DB, token))
        print(f"    {len(biz_rows)} business · {len(per_rows)} personal")

        dad_id, dad_info = find_dad_db(token)
        if dad_id:
            dad_rows = flatten(notion_query(dad_id, token))
            print(f"    {len(dad_rows)} appa ({dad_info})")
        else:
            dad_status = dad_info
            print(f"    appa's sheet: {dad_info}")
    else:
        print("  ! No .notion_token found — building with revenue only.")
        print("    See the setup notes at the top of this file to enable expenses.")

    expenses, cash, cash_cat, true_cost, delivery, vendors = shape_expenses(biz_rows)
    personal, personal_m, personal_cats = shape_personal(per_rows)
    dad = shape_shared(dad_rows)
    monthly = build_monthly(rev, cash, true_cost, delivery, personal_m)

    if not monthly:
        raise SystemExit("  No data to build from.")

    # Feature the latest month with REAL activity: revenue, meals delivered, or money
    # actually spent. Note the test is cash_out (actual purchases), NOT cogs_true —
    # a month that exists only because a bulk buy was amortised forward has cash_out
    # of zero and must never be featured. Deliberately not gated on today's date, so
    # a purchase logged for next month still shows up.
    today = f"{MONTH_ORDER[datetime.now().month - 1]}-{datetime.now().year}"
    real = [m for m in monthly
            if m["total_revenue"] or m["cash_out"] or m["delivered"]]
    cur = real[-1]["month"] if real else today

    os.makedirs(DATA_DIR, exist_ok=True)
    with open(os.path.join(DATA_DIR, "finance.json"), "w", encoding="utf-8") as f:
        json.dump({"generated": datetime.now().isoformat(), "monthly": monthly,
                   "expenses": expenses, "personal": personal, "dad": dad},
                  f, indent=1, default=str)

    print("  Building dashboard ...")
    with open(OUT_HTML, "w", encoding="utf-8") as f:
        f.write(render_dashboard(monthly, expenses, personal, dad, cur, dad_status))

    print("  Building Excel ...")
    ok = write_excel(monthly, expenses, personal, OUT_XLSX)

    print("  " + "-" * 44)
    print(f"  Dashboard : {OUT_HTML}")
    if ok:
        print(f"  Excel     : {OUT_XLSX}")
    print(f"  Featuring : {cur}\n")


if __name__ == "__main__":
    main()
